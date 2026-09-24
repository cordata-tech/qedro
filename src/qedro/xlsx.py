"""The Art. 30 record as a workbook.

The one output an auditor actually asks for, and the only one that is not text.
A DPO receives a spreadsheet, forwards it, and somebody eventually sorts and
filters it — which changes what this format has to do, not only how it looks.

**The verdict goes where the file opens.** A record that does not claim to be a
proof says so on the first sheet, above the table. Parking that behind the
activities, or on a tab nobody clicks, would reproduce the exact failure this
tool exists to prevent: an artefact whose holes are invisible. The Scope sheet
carries the reasons; the first sheet carries the verdict and points at them.

**Provenance is a column of words, never a colour.** The tint on an asserted
cell is reinforcement for a reader who can see it and carries nothing on its
own — every tinted cell also says ``mapping file`` or ``not declared`` in text.

Timestamps are written as ISO-8601 strings rather than as spreadsheet dates.
Excel's date type has no timezone and openpyxl refuses an aware datetime, so a
real date cell would mean dropping the offset from evidence whose whole value
is being exact about when something ran.
"""

from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from . import TOMBSTONE, words
from .deployer import SIX_MONTHS, DeployerRecord
from .errors import QedroError
from .provenance import Record as ProvenanceRecord
from .quality import Record as QualityRecord
from .ropa import Activity, Record, Sourced

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError as exc:
    # A declared dependency, so this means a broken install rather than a
    # choice — but a sentence still beats a traceback.
    raise QedroError(
        "xlsx output needs openpyxl, which installs with qedro — try "
        "`pip install openpyxl`, or ask for --format markdown"
    ) from exc

#: Column headings and the width each needs to be readable without the reader
#: dragging borders around.
COLUMNS: tuple[tuple[str, int], ...] = (
    ("Activity", 38),
    ("Domain", 14),
    ("Purpose", 24),
    ("Purpose source", 20),
    ("Legal basis", 22),
    ("Legal basis source", 20),
    ("Categories", 26),
    ("Special category", 22),
    ("Subjects", 18),
    ("Residency", 14),
    ("Retention", 14),
    ("Recipients", 30),
    ("Security measures", 30),
    ("Unclassified", 30),
    ("Reads", 34),
    ("Writes", 34),
    ("Runs", 7),
    ("First seen", 26),
    ("Last seen", 26),
    ("Notes", 46),
)

AT = {heading: index for index, (heading, _) in enumerate(COLUMNS, start=1)}
WRAPPED = {
    AT["Reads"],
    AT["Writes"],
    AT["Notes"],
    AT["Categories"],
    AT["Unclassified"],
    AT["Recipients"],
    AT["Security measures"],
}

#: How a provenance reads to somebody who did not write this tool.
SOURCE = {
    "facet": "emitted facet",
    "mapping": "mapping file",
    "declared": "declared, no lineage",
    "absent": "not declared",
}

#: Past this many dataset keys a cell stops being readable, and a cell has a
#: hard character limit besides. What was dropped is stated in the cell: a list
#: that simply stops looks complete.
DATASETS_PER_CELL = 40

#: The table starts here. Above it: the title, the controller, the contact and
#: the verdict.
HEADER_ROW = 6

TITLE = Font(bold=True, size=14)
STRONG = Font(bold=True)
HEADING = Font(bold=True, color="FFFFFFFF")
HEADING_FILL = PatternFill(fill_type="solid", start_color="FF1F2933", end_color="FF1F2933")
ASSERTED_FILL = PatternFill(fill_type="solid", start_color="FFFDF0D5", end_color="FFFDF0D5")
TOP = Alignment(vertical="top")
WRAP = Alignment(vertical="top", wrap_text=True)


def workbook(record: Record) -> bytes:
    """The record as ``.xlsx`` bytes, ready to be written or handed on."""
    book = Workbook()
    # Document properties, because a file that leaves a DPO's outbox should say
    # what made it without anyone having to open it.
    book.properties.creator = "qedro"
    book.properties.title = "Record of processing activities"
    book.properties.description = record.scope.OUT_OF_VIEW

    _activities(book.active, record)
    _scope(book.create_sheet("Scope"), record)

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _activities(sheet: Any, record: Record) -> None:
    sheet.title = "Art. 30 record"

    sheet["A1"] = "Record of processing activities"
    sheet["A1"].font = TITLE
    # Art. 30(1)(a). A blank here would read as a formatting slip rather than
    # as the hole it is.
    sheet["A2"] = record.controller.name or "No controller is declared"
    sheet["A2"].font = STRONG
    sheet["A3"] = f"Contact: {record.controller.contact}" if record.controller.contact else ""
    sheet["A4"] = _verdict(record)
    sheet["A4"].font = STRONG

    for index, (heading, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=heading)
        cell.font = HEADING
        cell.fill = HEADING_FILL
        cell.alignment = WRAP
        sheet.column_dimensions[get_column_letter(index)].width = width

    for offset, activity in enumerate(record.activities):
        _row(sheet, HEADER_ROW + 1 + offset, activity)
    for offset, entry in enumerate(record.declared, start=len(record.activities)):
        _declared_row(sheet, HEADER_ROW + 1 + offset, entry)

    last_column = get_column_letter(len(COLUMNS))
    last_row = HEADER_ROW + len(record.activities) + len(record.declared)
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{last_column}{last_row}"
    sheet.freeze_panes = f"A{HEADER_ROW + 1}"


def _row(sheet: Any, row: int, activity: Activity) -> None:
    values = (
        activity.key,
        activity.domain,
        activity.purpose.value or "—",
        _source(activity.purpose),
        activity.legal_basis.value or "—",
        _source(activity.legal_basis),
        *_classification(activity),
        _datasets(activity.inputs),
        _datasets(activity.outputs),
        activity.runs,
        _when(activity.first_seen),
        _when(activity.last_seen),
        "\n".join(activity.gaps()),
    )
    for index, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=index, value=value)
        cell.alignment = WRAP if index in WRAPPED else TOP

    # One tint, meaning *this cell is part of why the record does not claim to
    # be a proof*. Which of the two problems it is comes from the text in the
    # cell and from the Notes column; the colour never carries it alone. The
    # source cell is tinted when the value is not evidence, and the value cell
    # when the vocabulary does not define it — an emitted value can be perfectly
    # good evidence of a term nobody agreed on.
    for value_column, source_column, sourced in (
        (AT["Purpose"], AT["Purpose source"], activity.purpose),
        (AT["Legal basis"], AT["Legal basis source"], activity.legal_basis),
    ):
        if not sourced.evidenced:
            sheet.cell(row=row, column=source_column).fill = ASSERTED_FILL
        if sourced.unrecognised:
            sheet.cell(row=row, column=value_column).fill = ASSERTED_FILL

    # A classification outside a closed term is tinted where the value is, for
    # the same reason: the value is what nobody agreed on.
    for heading, key in CLASSIFICATION_COLUMNS:
        if any(c.unrecognised for c in activity.classification if c.key == key):
            sheet.cell(row=row, column=AT[heading]).fill = ASSERTED_FILL


#: The classification columns, in Art. 30(1) order, and the tag key each reads.
CLASSIFICATION_COLUMNS = (
    ("Categories", "data_category"),
    ("Special category", "special_category"),
    ("Subjects", "subject_type"),
    ("Residency", "residency"),
    ("Retention", "retention"),
)


def _classification(activity: Activity) -> tuple[Any, ...]:
    """The five classification cells, and what the activity did not classify.

    An empty cell is `—`, meaning nothing said, and the Unclassified column
    counts the datasets that carried no classification the record can use —
    because a row of dashes must not read as *no personal data here*.
    """
    cells: list[Any] = []
    for _, key in CLASSIFICATION_COLUMNS:
        values = [c for c in activity.classification if c.key == key]
        cells.append("\n".join(c.value for c in values) if values else "—")
    # Art. 30(1)(d) and (g), always marked: nothing emits either, so an
    # unmarked value here would read like the evidenced columns beside it.
    for asserted in (activity.recipients, activity.security_measures):
        cells.append("\n".join(asserted) + "\n(declared)" if asserted else "—")
    if activity.unclassified:
        total = len(set(activity.inputs) | set(activity.outputs))
        cells.append(
            f"{len(activity.unclassified)} of {total}:\n" + _datasets(activity.unclassified)
        )
    else:
        cells.append("—")
    return tuple(cells)


def _declared_row(sheet: Any, row: int, entry: Any) -> None:
    """A declared activity. Every cell the events would have filled says
    `no lineage` — an empty Reads cell reads as *touches no data*, and a 0 in
    Runs as *ran zero times*, and neither is known."""
    reads = "no lineage"
    if entry.inputs:
        reads = "no lineage — declared:\n" + "\n".join(entry.inputs)
    values = (
        entry.name,
        entry.domain,
        entry.purpose.value or "—",
        _source(entry.purpose),
        entry.legal_basis.value or "—",
        _source(entry.legal_basis),
        # A declared activity emits no lineage, so nothing classified its data
        # either. `no lineage` rather than `—`, which would read as *nothing
        # applies* instead of *nothing could have said*.
        *("no lineage",) * len(CLASSIFICATION_COLUMNS),
        "\n".join(entry.recipients) + "\n(declared)" if entry.recipients else "no lineage",
        "\n".join(entry.security_measures) + "\n(declared)"
        if entry.security_measures
        else "no lineage",
        "no lineage",
        reads,
        "no lineage",
        "no lineage",
        "no lineage",
        "no lineage",
        "\n".join(x for x in ("declared, no lineage", entry.note) if x),
    )
    for index, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=index, value=value)
        cell.alignment = WRAP if index in WRAPPED else TOP
    for column in (AT["Purpose source"], AT["Legal basis source"]):
        sheet.cell(row=row, column=column).fill = ASSERTED_FILL


def _verdict(record: Record) -> str:
    if record.complete:
        return f"Every activity in this record stands on emitted evidence. {TOMBSTONE}"
    n = len(record.completeness.reasons)
    return (
        f"This record does not claim to be a proof — {words.count(n, 'reason')} on the Scope sheet."
    )


def _source(sourced: Sourced) -> str:
    """Where the value came from, and whether the vocabulary knows it.

    The flag sits on the source cell rather than on the value, so the value
    column stays exactly what was declared and a reader can still filter on it.
    """
    label = SOURCE.get(str(sourced.provenance), str(sourced.provenance))
    return f"{label} — not in the vocabulary" if sourced.unrecognised else label


def _datasets(keys: tuple[str, ...]) -> str:
    if len(keys) <= DATASETS_PER_CELL:
        return "\n".join(keys)
    remaining = len(keys) - DATASETS_PER_CELL
    return "\n".join((*keys[:DATASETS_PER_CELL], f"… and {remaining} more"))


def _when(value: datetime | None) -> str:
    return value.isoformat() if value else ""


def _scope(sheet: Any, record: Record) -> None:
    scope = record.scope
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 96

    sheet["A1"] = "Scope of this record"
    sheet["A1"].font = TITLE

    # From `scope.lines()`, as the other projections' sheets already were. This
    # sheet listed its rows itself, which is how the declared count from #6 never
    # reached the workbook while every other format printed it.
    rows = [("Source", scope.source or "unknown"), ("Window", scope.window())]
    rows += [(label[0].upper() + label[1:], value) for label, value in scope.lines()]
    rows.append(("Vocabulary", record.vocabulary))
    if scope.domains_source():
        rows.append(("Domains", scope.domains_source()))
    if scope.domains_silent:
        rows.append(("Declared but silent", ", ".join(scope.domains_silent)))

    row = 3
    for label, value in rows:
        sheet.cell(row=row, column=1, value=label).font = STRONG
        sheet.cell(row=row, column=2, value=value).alignment = WRAP
        row += 1

    # Unconditional, on every run including one that earns the mark. See
    # cordata-tech/qedro#3.
    row += 1
    sheet.cell(row=row, column=1, value="Not covered").font = STRONG
    sheet.cell(row=row, column=2, value=scope.OUT_OF_VIEW).alignment = WRAP
    sheet.row_dimensions[row].height = 64

    row += 2
    sheet.cell(row=row, column=1, value="Completeness").font = STRONG
    if record.complete:
        sheet.cell(
            row=row,
            column=2,
            value=f"Every activity in this record stands on emitted evidence. {TOMBSTONE}",
        )
        return

    sheet.cell(row=row, column=2, value="This record does not claim to be a proof.").font = STRONG
    for reason in record.completeness.reasons:
        row += 1
        sheet.cell(row=row, column=2, value=reason).alignment = WRAP


# --- proving an erasure reached its descendants, #4 ------------------------


#: `Rewritten since` is the column a reader sorts by, and the blank cells are the
#: work list. There is deliberately no *erased* column: this reports datasets and
#: not rows, and a column headed *erased* would be read as the row-level claim the
#: whole projection refuses to make.
ERASURE_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Descendant", 44),
    ("Depth", 8),
    ("Derived by", 32),
    ("Rewritten since the erasure", 28),
    ("Runs since", 12),
    ("Unfinished runs", 16),
    ("State", 40),
)


def erasure_workbook(record: Any) -> bytes:
    """An erasure record as ``.xlsx`` bytes."""
    book = Workbook()
    book.properties.creator = "qedro"
    book.properties.title = f"Erasure of {record.dataset}"
    book.properties.description = record.scope.OUT_OF_VIEW

    _descendants(book.active, record)
    _erasure_scope(book.create_sheet("Scope"), record)

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _descendants(sheet: Any, record: Any) -> None:
    sheet.title = "Descendants"

    sheet["A1"] = f"Erasure of {record.dataset}"
    sheet["A1"].font = TITLE
    sheet["A2"] = f"Tombstone: {record.tombstone.describe()}"
    # The verdict goes where the file opens, as in every other workbook here.
    sheet["A3"] = _erasure_verdict(record)
    sheet["A3"].font = STRONG

    header = 5
    for index, (heading, width) in enumerate(ERASURE_COLUMNS, start=1):
        cell = sheet.cell(row=header, column=index, value=heading)
        cell.font = HEADING
        cell.fill = HEADING_FILL
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header + 1, column=1)

    for offset, d in enumerate(record.descendants, start=header + 1):
        values = (
            d.dataset,
            d.depth,
            d.through,
            _when(d.rewritten),
            d.runs_since,
            d.incomplete,
            _erasure_state(d),
        )
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=offset, column=index, value=value)
            cell.alignment = TOP
            # Tinted where the proof does not reach, and the State column says
            # the same thing in words — the tint carries nothing on its own.
            if not d.reached:
                cell.fill = ASSERTED_FILL


def _erasure_state(descendant: Any) -> str:
    if descendant.rewritten is not None:
        return "rewritten since the erasure"
    if descendant.incomplete:
        return "not known — a run started after the erasure and never finished"
    return "not rewritten since the erasure"


def _erasure_verdict(record: Any) -> str:
    if record.complete:
        return f"Every descendant was rewritten after an emitted erasure. {TOMBSTONE}"
    return "This record does not claim to be a proof — see the Scope sheet."


def _erasure_scope(sheet: Any, record: Any) -> None:
    scope = record.scope
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 96

    sheet["A1"] = "Scope of this record"
    sheet["A1"].font = TITLE

    rows = [("Source", scope.source or "unknown"), ("Window", scope.window())]
    rows += [(label[0].upper() + label[1:], value) for label, value in scope.lines()]

    row = 3
    for label, value in rows:
        sheet.cell(row=row, column=1, value=label).font = STRONG
        sheet.cell(row=row, column=2, value=value).alignment = WRAP
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Not covered").font = STRONG
    sheet.cell(row=row, column=2, value=scope.OUT_OF_VIEW).alignment = WRAP
    sheet.row_dimensions[row].height = 80

    row += 2
    sheet.cell(row=row, column=1, value="Completeness").font = STRONG
    if record.complete:
        sheet.cell(row=row, column=2, value=_erasure_verdict(record))
        return
    sheet.cell(row=row, column=2, value="This record does not claim to be a proof.").font = STRONG
    for reason in record.completeness.reasons:
        row += 1
        sheet.cell(row=row, column=2, value=reason).alignment = WRAP


# --- what changed between two records, #14 ---------------------------------


#: The findings table. `Before` and `After` are split from what each rests on,
#: because the column a reader sorts by to find lost evidence is the source, not
#: the value — and the value is often identical on both sides.
COMPARISON_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Activity", 38),
    ("Field", 22),
    ("Change", 14),
    ("Before", 28),
    ("Before source", 20),
    ("After", 28),
    ("After source", 20),
    ("Loses evidence", 16),
)

#: The same table when one side is a register. `Owner` is what turns a list of
#: disagreements into a list of things somebody can go and do, and it is the
#: column a reader sorts by. There is no `Register source` column: everything a
#: register says is an assertion, so a column repeating that on every row would
#: carry nothing.
DRIFT_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Activity", 38),
    ("Owner", 28),
    ("Field", 22),
    ("Finding", 14),
    ("Register says", 28),
    ("Record says", 28),
    ("Record's source", 20),
    ("What it means", 44),
)


def comparison_workbook(comparison: Any) -> bytes:
    """A comparison as ``.xlsx`` bytes.

    **No verdict cell of its own.** Every other workbook here opens with whether
    the artefact claims to be a proof; a comparison has no such claim to make, so
    what sits above the table is what the reader has to hold in mind instead —
    the cautions, if the two records do not cover the same ground.
    """
    book = Workbook()
    book.properties.creator = "qedro"
    book.properties.title = (
        "Where the register and the record disagree"
        if comparison.drift
        else "What changed between two records"
    )
    book.properties.description = comparison.scope.OUT_OF_VIEW

    _findings(book.active, comparison)
    _comparison_scope(book.create_sheet("Scope"), comparison)

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _findings(sheet: Any, comparison: Any) -> None:
    drift = comparison.drift
    sheet.title = "Findings"

    sheet["A1"] = (
        "Where the register and the record disagree"
        if drift
        else "What changed between two records"
    )
    sheet["A1"].font = TITLE
    joiner = " against " if drift else " → "
    sheet["A2"] = f"{comparison.before.origin}{joiner}{comparison.after.origin}"

    row = 3
    for caution in comparison.comparability.cautions():
        sheet.cell(row=row, column=1, value=f"! {caution}").font = STRONG
        row += 1

    columns = DRIFT_COLUMNS if drift else COMPARISON_COLUMNS
    header = row + 1
    for index, (heading, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=header, column=index, value=heading)
        cell.font = HEADING
        cell.fill = HEADING_FILL
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header + 1, column=1)

    for offset, change in enumerate(comparison.changes, start=header + 1):
        values = _drift_row(change) if drift else _comparison_row(change)
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=offset, column=index, value=value)
            cell.alignment = TOP
            # No tint on a drift row. Every one of them is a disagreement, and a
            # sheet tinted end to end says nothing that the rows do not.
            if change.regression and not drift:
                cell.fill = ASSERTED_FILL


def _comparison_row(change: Any) -> tuple[str, ...]:
    return (
        change.entry,
        change.label,
        change.kind,
        _comparison_value(change.before),
        _comparison_source(change.before),
        _comparison_value(change.after),
        _comparison_source(change.after),
        # A word rather than TRUE, because this column is the one a reader
        # filters on and `TRUE` says nothing about what was lost.
        "evidence lost" if change.regression else "",
    )


def _drift_row(change: Any) -> tuple[str, ...]:
    """One disagreement between a register and a record.

    What the record's side rests on gets its own column: a register contradicting
    an emitted facet is a different problem from two hand-maintained documents
    disagreeing, and a reader deciding which rows to chase needs to sort by it.
    """
    from .render import DISAGREEMENT

    means = ""
    if change.before is not None and change.after is not None:
        means = DISAGREEMENT.get(str(change.after.provenance), "")
    elif change.kind == "removed":
        means = "no pipeline emitted it in the window"
    elif change.kind == "added":
        means = "nobody wrote this one down"
    return (
        change.entry,
        change.owner,
        change.label,
        change.kind,
        _comparison_value(change.before),
        _comparison_value(change.after),
        _comparison_source(change.after),
        means,
    )


def _comparison_value(value: Any) -> str:
    """The value, or a word for there having been no side at all."""
    if value is None:
        return "—"
    return value.value or "not stated"


def _comparison_source(value: Any) -> str:
    if value is None:
        return "—"
    return SOURCE.get(str(value.provenance), str(value.provenance))


def _comparison_scope(sheet: Any, comparison: Any) -> None:
    scope = comparison.scope
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 96

    sheet["A1"] = "Scope of this comparison"
    sheet["A1"].font = TITLE

    rows = [("Source", scope.source or "unknown"), ("Window", scope.window())]
    rows += [(label[0].upper() + label[1:], value) for label, value in scope.lines()]

    row = 3
    for label, value in rows:
        sheet.cell(row=row, column=1, value=label).font = STRONG
        sheet.cell(row=row, column=2, value=value).alignment = WRAP
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Not covered").font = STRONG
    sheet.cell(row=row, column=2, value=scope.OUT_OF_VIEW).alignment = WRAP
    sheet.row_dimensions[row].height = 64

    # Each record's verdict, never this comparison's. A workbook that printed a
    # mark here would be claiming something no comparison can earn.
    row += 2
    heading = "The documents' own verdicts" if comparison.drift else "The records' own verdicts"
    sheet.cell(row=row, column=1, value=heading).font = STRONG
    sides = (
        (("register", comparison.register), ("record", comparison.record))
        if comparison.drift
        else (("before", comparison.before), ("after", comparison.after))
    )
    for label, side in sides:
        if side.is_register:
            sheet.cell(
                row=row,
                column=2,
                value=f"{label}: a hand-maintained document; it makes no claim",
            )
            row += 1
            continue
        if side.complete:
            sheet.cell(row=row, column=2, value=f"{label}: stands on its own evidence")
            row += 1
            continue
        sheet.cell(row=row, column=2, value=f"{label}: does not claim to be a proof")
        for reason in side.reasons:
            row += 1
            sheet.cell(row=row, column=2, value=reason).alignment = WRAP
        row += 1


# --- the assertion history -------------------------------------------------

#: One row per expectation rather than per dataset. A reader filtering for
#: `FAILED` wants the expectation that failed, not the table it was on.
QUALITY_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Dataset", 44),
    ("Domain", 14),
    ("Expectation", 40),
    ("Column", 20),
    ("Result", 12),
    ("Runs", 8),
    ("Failed", 8),
    ("Last failure", 26),
    ("Last asserted", 26),
    ("Asserted by", 34),
)

QUALITY_AT = {heading: index for index, (heading, _) in enumerate(QUALITY_COLUMNS, start=1)}


def quality_workbook(record: QualityRecord) -> bytes:
    """The assertion history as ``.xlsx`` bytes.

    Three sheets, and the middle one is the point. *Not checked* is a sheet of
    its own rather than a note at the bottom of the first, because a reader who
    stops after the green table has been misled — and a tab they have to
    dismiss is harder to stop at than a paragraph they can skim past.
    """
    book = Workbook()
    book.properties.creator = "qedro"
    book.properties.title = "Data quality assertion history"
    book.properties.description = record.scope.OUT_OF_VIEW

    _expectations(book.active, record)
    _unchecked(book.create_sheet("Not checked"), record)
    _quality_scope(book.create_sheet("Scope"), record)

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _expectations(sheet: Any, record: QualityRecord) -> None:
    sheet.title = "Assertions"

    sheet["A1"] = "Data quality assertion history"
    sheet["A1"].font = TITLE
    sheet["A2"] = record.controller.name or "No controller is declared"
    sheet["A2"].font = STRONG
    sheet["A3"] = (
        f"{record.scope.checked} datasets checked, "
        f"{record.scope.unchecked} with no assertions at all"
    )
    sheet["A4"] = _quality_verdict(record)
    sheet["A4"].font = STRONG

    for index, (heading, width) in enumerate(QUALITY_COLUMNS, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=heading)
        cell.font = HEADING
        cell.fill = HEADING_FILL
        cell.alignment = WRAP
        sheet.column_dimensions[get_column_letter(index)].width = width

    row = HEADER_ROW
    for dataset in record.datasets:
        for expectation in dataset.expectations:
            row += 1
            values = (
                dataset.key,
                dataset.domain,
                expectation.assertion,
                expectation.column,
                "held" if expectation.holds else "FAILED",
                expectation.runs,
                expectation.failures,
                _when(expectation.last_failure),
                _when(expectation.last_seen),
                ", ".join(dataset.asserted_by),
            )
            for index, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=index, value=value)
                cell.alignment = TOP
            if not expectation.holds:
                sheet.cell(row=row, column=QUALITY_AT["Result"]).fill = ASSERTED_FILL

    sheet.auto_filter.ref = (
        f"A{HEADER_ROW}:{get_column_letter(len(QUALITY_COLUMNS))}{max(row, HEADER_ROW)}"
    )
    sheet.freeze_panes = f"A{HEADER_ROW + 1}"


def _quality_verdict(record: QualityRecord) -> str:
    if record.complete:
        return f"This history accounts for every dataset in view. {TOMBSTONE}"
    n = len(record.completeness.reasons)
    return (
        f"This history does not claim to be complete — "
        f"{words.count(n, 'reason')} on the Scope sheet."
    )


def _unchecked(sheet: Any, record: QualityRecord) -> None:
    sheet.column_dimensions["A"].width = 60

    sheet["A1"] = "Datasets with no assertions"
    sheet["A1"].font = TITLE
    sheet["A2"] = (
        "These appeared in lineage in the window and nothing asserted anything "
        "about them. Not checked is not the same as passed."
    )
    sheet["A2"].alignment = WRAP
    sheet.row_dimensions[2].height = 32

    # Every one of them, not a sample: this is the sheet somebody works
    # through, and a truncated work list is a finished work list.
    for offset, key in enumerate(record.unchecked):
        sheet.cell(row=4 + offset, column=1, value=key).alignment = TOP


def _quality_scope(sheet: Any, record: QualityRecord) -> None:
    scope = record.scope
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 96

    sheet["A1"] = "Scope of this history"
    sheet["A1"].font = TITLE

    rows = [("Source", scope.source or "unknown"), ("Window", scope.window())]
    rows += list(scope.lines())
    if scope.domains_source():
        rows.append(("Domains", scope.domains_source()))
    if scope.domains_silent:
        rows.append(("Declared but silent", ", ".join(scope.domains_silent)))

    row = 3
    for label, value in rows:
        sheet.cell(row=row, column=1, value=label[0].upper() + label[1:]).font = STRONG
        sheet.cell(row=row, column=2, value=value).alignment = WRAP
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Not covered").font = STRONG
    sheet.cell(row=row, column=2, value=scope.OUT_OF_VIEW).alignment = WRAP
    sheet.row_dimensions[row].height = 64

    row += 2
    sheet.cell(row=row, column=1, value="Completeness").font = STRONG
    if record.complete:
        sheet.cell(row=row, column=2, value=_quality_verdict(record))
        return

    sheet.cell(row=row, column=2, value="This history does not claim to be complete.").font = STRONG
    for reason in record.completeness.reasons:
        row += 1
        sheet.cell(row=row, column=2, value=reason).alignment = WRAP


# --- the provenance chain ---------------------------------------------------

PROVENANCE_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Depth", 8),
    ("Dataset", 44),
    ("Produced by", 34),
    ("Run", 38),
    ("When", 26),
    ("Repository", 40),
    ("Commit", 16),
    ("Branch", 14),
    ("Path", 34),
    ("Signature", 34),
    # One hop further out than the commit: the application release behind
    # the data the run read, rather than behind the pipeline itself.
    ("Source published by", 40),
    ("Chain ends", 44),
)

PROVENANCE_AT = {h: i for i, (h, _) in enumerate(PROVENANCE_COLUMNS, start=1)}


def provenance_workbook(record: ProvenanceRecord) -> bytes:
    """The chain as ``.xlsx`` bytes.

    One sheet, because a chain is one thing. Depth is a column rather than
    indentation: a reader who sorts the sheet would destroy indentation without
    noticing, and the depth is the part they would be sorting by.
    """
    book = Workbook()
    book.properties.creator = "qedro"
    book.properties.title = f"Provenance of {record.dataset}"
    book.properties.description = record.scope.OUT_OF_VIEW

    _chain(book.active, record)
    _provenance_scope(book.create_sheet("Scope"), record)

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _chain(sheet: Any, record: ProvenanceRecord) -> None:
    sheet.title = "Chain"

    sheet["A1"] = f"Provenance of {record.dataset}"
    sheet["A1"].font = TITLE
    sheet["A2"] = record.controller.name or "No controller is declared"
    sheet["A2"].font = STRONG
    sheet["A3"] = (
        f"{record.scope.with_commit} of {record.scope.steps} steps name a commit, "
        f"{record.scope.with_signature} report a signature, "
        f"{record.scope.with_publisher} name the application that published what they read"
    )
    sheet["A4"] = _provenance_verdict(record)
    sheet["A4"].font = STRONG

    for index, (heading, width) in enumerate(PROVENANCE_COLUMNS, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=heading)
        cell.font = HEADING
        cell.fill = HEADING_FILL
        cell.alignment = WRAP
        sheet.column_dimensions[get_column_letter(index)].width = width

    for offset, step in enumerate(record.steps):
        row = HEADER_ROW + 1 + offset
        p = step.production
        values = (
            step.depth,
            step.dataset,
            p.job if p else "",
            p.run_id if p else "",
            _when(p.when) if p else "",
            p.code.repository if p else "",
            p.code.short if p else "",
            p.code.branch if p else "",
            p.code.path if p else "",
            p.signature.describe() if p else "",
            p.published.describe() if p else "",
            step.ended,
        )
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=index, value=value)
            cell.alignment = WRAP if index == PROVENANCE_AT["Path"] else TOP

        # Tinted where the chain cannot show authorisation — and the cell says
        # which of the three reasons it is, so the colour carries nothing alone.
        if not p or not p.authorised:
            sheet.cell(row=row, column=PROVENANCE_AT["Signature"]).fill = ASSERTED_FILL
        if step.ended:
            sheet.cell(row=row, column=PROVENANCE_AT["Chain ends"]).fill = ASSERTED_FILL

    last = HEADER_ROW + len(record.steps)
    sheet.auto_filter.ref = (
        f"A{HEADER_ROW}:{get_column_letter(len(PROVENANCE_COLUMNS))}{max(last, HEADER_ROW)}"
    )
    sheet.freeze_panes = f"A{HEADER_ROW + 1}"


def _provenance_verdict(record: ProvenanceRecord) -> str:
    if record.complete:
        return f"Every step in this chain ran under a signed commit. {TOMBSTONE}"
    n = len(record.completeness.reasons)
    return (
        f"This chain does not claim to be a proof — {words.count(n, 'reason')} on the Scope sheet."
    )


def _provenance_scope(sheet: Any, record: ProvenanceRecord) -> None:
    scope = record.scope
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 96

    sheet["A1"] = "Scope of this chain"
    sheet["A1"].font = TITLE

    rows = [("Source", scope.source or "unknown"), ("Window", scope.window())]
    rows += list(scope.lines())

    row = 3
    for label, value in rows:
        sheet.cell(row=row, column=1, value=label[0].upper() + label[1:]).font = STRONG
        sheet.cell(row=row, column=2, value=value).alignment = WRAP
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Not covered").font = STRONG
    sheet.cell(row=row, column=2, value=scope.OUT_OF_VIEW).alignment = WRAP
    sheet.row_dimensions[row].height = 64

    row += 2
    sheet.cell(row=row, column=1, value="Completeness").font = STRONG
    if record.complete:
        sheet.cell(row=row, column=2, value=_provenance_verdict(record))
        return

    sheet.cell(row=row, column=2, value="This chain does not claim to be a proof.").font = STRONG
    for reason in record.completeness.reasons:
        row += 1
        sheet.cell(row=row, column=2, value=reason).alignment = WRAP


# --- the deployer view ------------------------------------------------------

#: One row per use case, lineage and declared alike, so a reader filtering on
#: `Evidence` sees both kinds side by side instead of on separate sheets.
DEPLOYER_COLUMNS: tuple[tuple[str, int], ...] = (
    ("Use case", 38),
    ("Evidence", 22),
    ("Domain", 12),
    ("Purpose", 24),
    ("Purpose source", 20),
    ("Legal basis", 22),
    ("Legal basis source", 20),
    ("Model version", 34),
    ("Latest run", 38),
    ("Inputs read, latest run (context for Art. 26(4))", 40),
    ("Run records in view (context for Art. 26(6))", 44),
    ("Notes", 46),
)

DEPLOYER_AT = {h: i for i, (h, _) in enumerate(DEPLOYER_COLUMNS, start=1)}
DEPLOYER_WRAPPED = {
    DEPLOYER_AT["Model version"],
    DEPLOYER_AT["Inputs read, latest run (context for Art. 26(4))"],
    DEPLOYER_AT["Run records in view (context for Art. 26(6))"],
    DEPLOYER_AT["Notes"],
}


def deployer_workbook(record: DeployerRecord) -> bytes:
    """The deployer view as ``.xlsx`` bytes.

    Row 3 carries the legal reference and its condition, above the table, for
    the same reason the verdict does: a sheet read without it would imply that
    Art. 26 applies today.
    """
    book = Workbook()
    book.properties.creator = "qedro"
    book.properties.title = "AI use cases, deployer view of the Art. 30 record"
    book.properties.description = record.scope.OUT_OF_VIEW

    _use_cases(book.active, record)
    _deployer_scope(book.create_sheet("Scope"), record)

    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def _use_cases(sheet: Any, record: DeployerRecord) -> None:
    sheet.title = "AI use cases"

    sheet["A1"] = "AI use cases, deployer view of the Art. 30 record"
    sheet["A1"].font = TITLE
    sheet["A2"] = record.controller.name or "No controller is declared"
    sheet["A2"].font = STRONG
    sheet["A3"] = f"{record.regulation}. {record.application}"
    sheet["A4"] = _deployer_verdict(record)
    sheet["A4"].font = STRONG

    for index, (heading, width) in enumerate(DEPLOYER_COLUMNS, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=index, value=heading)
        cell.font = HEADING
        cell.fill = HEADING_FILL
        cell.alignment = WRAP
        sheet.column_dimensions[get_column_letter(index)].width = width

    row = HEADER_ROW
    for use_case in record.use_cases:
        row += 1
        a = use_case.activity
        first, last, days = a.first_seen, a.last_seen, use_case.span_days
        if days is None or first is None or last is None:
            records = "no dated run records in view"
        else:
            enough = "at least six months" if days >= SIX_MONTHS else "less than six months"
            records = (
                f"{words.count(a.runs, 'run')}, {first.date()} to {last.date()}, "
                f"{words.count(days, 'day')} — {enough} of records in view"
            )
        versions = [
            f"{m.version} ({words.count(m.runs, 'run')}, reported by {m.reported_by})"
            for m in use_case.models
        ]
        if use_case.runs_without_model:
            versions.append(f"none reported ({words.count(use_case.runs_without_model, 'run')})")
        latest = use_case.latest
        values = (
            a.key,
            "emitted lineage",
            a.domain,
            a.purpose.value or "—",
            _source(a.purpose),
            a.legal_basis.value or "—",
            _source(a.legal_basis),
            "\n".join(versions),
            f"{latest.run_id}\n{_when(latest.when)}",
            "\n".join(latest.inputs),
            records,
            "\n".join(a.gaps()),
        )
        _deployer_row(sheet, row, values)
        for value_column, source_column, sourced in (
            (DEPLOYER_AT["Purpose"], DEPLOYER_AT["Purpose source"], a.purpose),
            (DEPLOYER_AT["Legal basis"], DEPLOYER_AT["Legal basis source"], a.legal_basis),
        ):
            if not sourced.evidenced:
                sheet.cell(row=row, column=source_column).fill = ASSERTED_FILL
            if sourced.unrecognised:
                sheet.cell(row=row, column=value_column).fill = ASSERTED_FILL

    for entry in record.declared:
        row += 1
        values = (
            entry.name,
            "declared, no lineage",
            entry.domain,
            entry.purpose.value or "—",
            _source(entry.purpose),
            entry.legal_basis.value or "—",
            _source(entry.legal_basis),
            f"{entry.model} (declared)" if entry.model else "not declared",
            "none",
            "\n".join(f"{i} (declared)" for i in entry.inputs) or "not declared",
            "none — nothing emits lineage for this use",
            entry.note,
        )
        _deployer_row(sheet, row, values)
        # The whole row is an assertion, so the cell that says so is tinted —
        # and says so in words, which is what carries it.
        sheet.cell(row=row, column=DEPLOYER_AT["Evidence"]).fill = ASSERTED_FILL

    last_column = get_column_letter(len(DEPLOYER_COLUMNS))
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{last_column}{max(row, HEADER_ROW)}"
    sheet.freeze_panes = f"A{HEADER_ROW + 1}"


def _deployer_row(sheet: Any, row: int, values: tuple) -> None:
    for index, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=index, value=value)
        cell.alignment = WRAP if index in DEPLOYER_WRAPPED else TOP


def _deployer_verdict(record: DeployerRecord) -> str:
    if record.complete:
        return f"Every AI use case in this view stands on emitted evidence. {TOMBSTONE}"
    n = len(record.completeness.reasons)
    return (
        f"This view does not claim to be a proof — {words.count(n, 'reason')} on the Scope sheet."
    )


def _deployer_scope(sheet: Any, record: DeployerRecord) -> None:
    scope = record.scope
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 96

    sheet["A1"] = "Scope of this view"
    sheet["A1"].font = TITLE

    rows = [
        ("Law", f"{record.regulation}. {record.application}"),
        ("Source", scope.source or "unknown"),
        ("Window", scope.window()),
    ]
    rows += list(scope.lines())

    row = 3
    for label, value in rows:
        sheet.cell(row=row, column=1, value=label[0].upper() + label[1:]).font = STRONG
        sheet.cell(row=row, column=2, value=value).alignment = WRAP
        row += 1

    row += 1
    sheet.cell(row=row, column=1, value="Not covered").font = STRONG
    sheet.cell(row=row, column=2, value=scope.OUT_OF_VIEW).alignment = WRAP
    sheet.row_dimensions[row].height = 96

    row += 2
    sheet.cell(row=row, column=1, value="Completeness").font = STRONG
    if record.complete:
        sheet.cell(row=row, column=2, value=_deployer_verdict(record))
        return

    sheet.cell(row=row, column=2, value="This view does not claim to be a proof.").font = STRONG
    for reason in record.completeness.reasons:
        row += 1
        sheet.cell(row=row, column=2, value=reason).alignment = WRAP
