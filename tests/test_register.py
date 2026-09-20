"""A hand-maintained register, checked against the record — cordata-tech/qedro#24.

Two things carry the weight. **A register is never merged into a record**, so
nothing here can assert itself into the document it is being compared with. And
**a field the register does not carry is not a disagreement**: a register is a
partial document by nature, and reporting everything it omits would bury the
handful of findings that matter under a page of things nobody claimed.
"""

from __future__ import annotations

import json as json_lib

import pytest

from qedro import compare, register, render
from qedro.errors import ConfigError, UsageError
from qedro.ropa import Provenance

from .test_compare import activity, document

ROW = """
register:
  fraud-scoring:
    job: acme.fraud/scored
    owner: Risk Analytics, T. Brandt
    purpose: fraud-detection
    legal_basis: legitimate-interest
"""


def demo_record():
    """The demo record as JSON, through the CLI, so the test reads what ships."""
    import subprocess
    import sys

    out = subprocess.run(
        [
            sys.executable,
            "-m",
            "qedro",
            "ropa",
            "demo/lineage-declared",
            "--config",
            "demo/qedro.yaml",
            "--format",
            "json",
        ],
        capture_output=True,
        text=True,
        check=True,
    )
    return json_lib.loads(out.stdout)


def demo_drift():
    return compare.compare(
        compare.read("demo/register.yaml"),
        compare.side(demo_record(), origin="record.json"),
    )


def register_side(text=ROW, *, origin="register.yaml"):
    from qedro import document as document_module

    parsed = document_module.parse(text, origin=origin)
    assert parsed is not None
    return compare.register_side(register.parse(parsed, origin=origin), origin=origin)


def drift(text=ROW, activities=None):
    return compare.compare(
        register_side(text),
        compare.side(
            document(activities if activities is not None else [activity()]), origin="record.json"
        ),
    )


class TestReadingTheDocument:
    def test_an_entry_naming_a_job_is_keyed_by_it(self):
        side = register_side()
        assert list(side.entries) == ["acme.fraud/scored"]

    def test_an_entry_with_no_job_is_keyed_by_the_name_it_was_written_under(self):
        side = register_side("register:\n  payroll-run:\n    purpose: payroll\n")
        assert list(side.entries) == ["payroll-run"]

    def test_the_name_it_was_written_under_is_kept_for_the_reader(self):
        # A row headed `fraud-scoring` that names a job is filed under the job
        # and has to be reported by the name its owner will recognise.
        entry = register_side().entries["acme.fraud/scored"]
        assert entry.name == "fraud-scoring"
        assert entry.owner == "Risk Analytics, T. Brandt"

    def test_everything_a_register_says_is_declared(self):
        entry = register_side().entries["acme.fraud/scored"]
        assert entry.single["purpose"].provenance is Provenance.DECLARED
        assert entry.single["legal basis"].provenance is Provenance.DECLARED

    def test_inputs_are_compared_against_the_record_s_reads(self):
        side = register_side(ROW + "    inputs:\n      - wh/raw.customers\n")
        entry = side.entries["acme.fraud/scored"]
        assert [v.value for v in entry.sets["reads"]] == ["wh/raw.customers"]

    def test_an_unknown_key_is_a_typo_until_proven_otherwise(self):
        with pytest.raises(ConfigError, match="unknown keys sensitivity"):
            register_side("register:\n  a:\n    sensitivity: high\n")

    def test_two_rows_for_one_activity_are_refused(self):
        text = (
            "register:\n"
            "  one:\n    job: acme.fraud/scored\n    purpose: a\n"
            "  two:\n    job: acme.fraud/scored\n    purpose: b\n"
        )
        with pytest.raises(ConfigError, match="two entries for"):
            register_side(text)

    def test_a_document_with_no_register_mapping_says_so(self):
        with pytest.raises(ConfigError, match="no `register:` mapping"):
            register_side("activities:\n  a:\n    purpose: p\n")

    def test_it_is_recognised_by_its_key_not_its_suffix(self, tmp_path):
        # A register may be written as JSON and a record always is, so the
        # suffix says nothing while the document says it plainly.
        path = tmp_path / "register.json"
        path.write_text(json_lib.dumps({"register": {"a": {"purpose": "payroll"}}}))
        assert compare.read(str(path)).is_register


class TestWhatIsAndIsNotADisagreement:
    def test_a_field_the_register_does_not_carry_is_not_compared(self):
        """A register says nothing about residency or retention, and reporting
        every such field would bury the findings that matter."""
        result = drift()
        assert [c.label for c in result.changes] == []

    def test_a_value_both_sides_state_and_agree_on_is_not_a_finding(self):
        # The register is declared and the record evidenced by construction, so
        # a difference of provenance between them is what the two documents
        # *are* rather than something either got wrong.
        assert drift().changes == ()
        assert drift().unchanged == 1

    def test_a_contradicted_purpose_names_which_side_is_evidence(self):
        result = drift(ROW.replace("fraud-detection", "marketing"))
        [change] = result.changes
        assert change.field == "purpose"
        assert change.before.value == "marketing"
        assert change.after.provenance is Provenance.FACET

    def test_a_field_the_register_carries_and_the_record_cannot_see(self):
        result = drift(ROW + "    model: churn-v4\n")
        [change] = [c for c in result.changes if c.field == "model"]
        assert change.after.value == ""

    def test_an_entry_only_in_the_register(self):
        result = drift(ROW.replace("acme.fraud/scored", "acme.fraud/gone"))
        kinds = {c.entry: c.kind for c in result.changes}
        assert kinds["acme.fraud/gone"] == compare.REMOVED
        assert kinds["acme.fraud/scored"] == compare.ADDED

    def test_the_owner_rides_on_every_finding_about_that_row(self):
        result = drift(ROW.replace("fraud-detection", "marketing"))
        assert result.changes[0].owner == "Risk Analytics, T. Brandt"

    def test_nothing_in_a_drift_report_is_called_a_loss_of_evidence(self):
        """Nothing was lost — the two documents were never in step."""
        result = drift(ROW.replace("fraud-detection", "marketing"))
        payload = json_lib.loads(render.json(result))
        assert all(f["loses_evidence"] is False for f in payload["findings"])


class TestRefusals:
    def test_two_registers_are_refused(self):
        with pytest.raises(UsageError, match="both registers"):
            compare.compare(register_side(), register_side(origin="other.yaml"))

    def test_the_order_of_the_two_paths_does_not_matter(self):
        record = compare.side(document([activity()]), origin="record.json")
        one = compare.compare(register_side(), record)
        other = compare.compare(record, register_side())
        assert one.changes == other.changes
        assert one.register.origin == other.register.origin == "register.yaml"


class TestHowItReads:
    def test_the_two_sides_are_named_register_and_record(self):
        out = render.text(drift(ROW.replace("fraud-detection", "marketing")))
        assert "Where the register and the record disagree" in out
        assert "register      demo" in out or "register      register.yaml" in out
        assert "record        record.json" in out

    def test_a_register_has_no_verdict_to_report(self):
        out = render.text(drift())
        assert "a hand-maintained document; it makes no claim" in out

    def test_the_standing_sentence_is_the_comparison_s_own(self):
        out = render.text(drift())
        assert "a register is what its owners wrote down" in out

    def test_the_summary_counts_which_side_each_missing_entry_is_on(self):
        out = render.text(drift(ROW.replace("acme.fraud/scored", "acme.fraud/gone")))
        assert "1 in the register only, 1 in the record only" in out

    def test_agreement_is_reported_rather_than_an_empty_page(self):
        assert "every entry in the register agrees with the record" in render.text(drift())

    def test_a_model_the_art30_view_cannot_carry_is_explained_not_just_reported(self):
        # Otherwise every AI register row reads as drift against a record that
        # has no field for a model at all.
        result = drift(ROW + "    model: churn-v4\n")
        assert any("view deployer" in c for c in result.comparability.cautions())

    def test_the_workbook_has_an_owner_column(self):
        from io import BytesIO

        from openpyxl import load_workbook

        book = load_workbook(BytesIO(render.xlsx(drift(ROW.replace("fraud-detection", "x")))))
        headings = [c.value for row in book["Findings"].iter_rows() for c in row if c.value]
        assert "Owner" in headings
        assert "Register says" in headings


class TestTheDemoRegister:
    """The acceptance case: platform#48's claim, made checkable.

    `demo/register.yaml` is four rows and four ways a register drifts, none of
    them invented for the demo — each is a thing that happens to a document
    nobody runs. One of the four is still true, so the report contains agreement
    as well as drift; a report where everything disagrees teaches a reader to
    distrust the tool rather than the register.
    """

    def test_a_row_that_is_still_true_reports_nothing_at_all(self):
        result = demo_drift()
        assert "acme.fraud/transactions-scored-daily" not in dict(result.by_entry())
        assert result.unchanged == 1

    def test_a_purpose_the_pipeline_changed_and_the_register_did_not(self):
        purpose = next(
            c
            for c in demo_drift().changes
            if c.entry == "acme.crm/customers-curated" and c.field == "purpose"
        )
        assert (purpose.before.value, purpose.after.value) == (
            "marketing",
            "customer-administration",
        )
        # The record's side is evidence, so the register is what is wrong here.
        assert purpose.after.provenance is Provenance.FACET

    def test_a_row_for_a_pipeline_that_no_longer_exists_names_who_owns_it(self):
        [gone] = dict(demo_drift().by_entry())["acme.billing/dunning-legacy"]
        assert gone.kind == compare.REMOVED
        assert gone.owner == "Billing, unassigned since 2025-11"

    def test_a_model_nobody_can_evidence(self):
        """platform#48's refusal (b): the register names a model in production
        and no run ever reported one."""
        model = next(
            c
            for c in demo_drift().changes
            if c.entry == "acme.crm/consent-sync" and c.field == "model"
        )
        assert model.before.value == "churn-v4"
        assert model.after.value == ""

    def test_pipelines_the_register_never_heard_of_are_reported(self):
        added = {c.entry for c in demo_drift().changes if c.kind == compare.ADDED}
        assert "acme.billing/invoices-nightly" in added
        assert "acme.fraud/scores-validated" in added
