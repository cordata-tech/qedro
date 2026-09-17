"""The demo estate, and the acceptance criteria it exists to make runnable.

cordata-tech/qedro#2 asks for two things that cannot be checked without a
committed estate: that `qedro ropa ./demo` produces a valid Art. 30 record with
no Cordata infrastructure present, and that the same command against
runtime-declared events produces a *materially richer* one.

Both are asserted here rather than described. The three-tier progression —
lineage alone, lineage plus the mapping file, lineage that declares — is the
adoption story the README tells, so a change that quietly broke a tier would
otherwise only be caught by somebody reading the docs and trying it.
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from tools import seed

from qedro import TOMBSTONE
from qedro.__main__ import main

DEMO = Path("demo")
PLAIN = str(DEMO / "lineage")
DECLARED = str(DEMO / "lineage-declared")
CONFIG = str(DEMO / "qedro.yaml")


def record(*argv: str, capsys) -> dict:
    assert main(["ropa", *argv, "--format", "json"]) == 0
    return json.loads(capsys.readouterr().out)


class TestTheEstateIsThere:
    def test_both_estates_are_committed(self):
        # Without this the acceptance criteria in #2 are unrunnable, which is
        # the state they were in until the estate existed.
        for directory in (DEMO / "lineage", DEMO / "lineage-declared"):
            assert sorted(p.name for p in directory.glob("*.ndjson")) == [
                "airflow.ndjson",
                "dbt.ndjson",
                "spark.ndjson",
            ]

    def test_the_two_estates_differ_only_in_the_declaration(self):
        # The comparison in demo/README.md is only worth anything if the
        # estates are otherwise identical. Same jobs, same runs, same times.
        def skeleton(directory: Path) -> list[tuple[str, str, str]]:
            out = []
            for path in sorted(directory.glob("*.ndjson")):
                for line in path.read_text(encoding="utf-8").splitlines():
                    e = json.loads(line)
                    out.append((e["eventTime"], e["job"]["name"], e["run"]["runId"]))
            return sorted(out)

        assert skeleton(DEMO / "lineage") == skeleton(DEMO / "lineage-declared")

    def test_it_is_generated_rather_than_edited(self):
        # The same guard CI runs. A demo edited by hand drifts from the
        # generator and the next regeneration silently reverts the edit.
        for directory, files in seed.estates().items():
            for name, body in files.items():
                assert (DEMO / directory / name).read_text(encoding="utf-8") == body, (
                    f"demo/{directory}/{name} is out of date — run `python tools/seed.py`"
                )


class TestSourceNeutrality:
    """cordata-tech/qedro#2, and the first commitment in CLAUDE.md.

    The acceptance test is that the demo runs with *no AWS anywhere in the
    picture*. A demo estate that quietly carried LakeFormation tags would make
    the claim untestable while looking like it proved it.
    """

    @pytest.mark.parametrize("directory", ["lineage", "lineage-declared"])
    def test_nothing_in_the_estate_names_one_cloud(self, directory):
        forbidden = ("lf_tags", "lakeformation", "arn:aws", "amazonaws", "cordata://")
        for path in (DEMO / directory).glob("*.ndjson"):
            body = path.read_text(encoding="utf-8").lower()
            for term in forbidden:
                assert term not in body, f"{path} contains {term!r}"

    def test_the_emitters_are_the_ones_people_already_run(self):
        producers = set()
        for path in (DEMO / "lineage").glob("*.ndjson"):
            for line in path.read_text(encoding="utf-8").splitlines():
                producers.add(json.loads(line)["producer"])
        assert len(producers) == 3
        assert all("OpenLineage/OpenLineage" in p for p in producers)


class TestLineageAlone:
    """Tier one: what the tool can say from lineage and nothing else."""

    def test_it_finds_the_estate(self, capsys):
        payload = record(PLAIN, capsys=capsys)
        assert len(payload["activities"]) == 6
        assert payload["scope"]["events"] == 216

    def test_but_cannot_say_why_any_of_it_happened(self, capsys):
        payload = record(PLAIN, capsys=capsys)
        assert payload["scope"]["provenance"]["undeclared"] == 6
        assert payload["complete"] is False

    def test_and_says_that_the_controller_is_missing(self, capsys):
        payload = record(PLAIN, capsys=capsys)
        assert any("no controller is declared" in r for r in payload["reasons"])


class TestLineagePlusTheMappingFile:
    """Tier two: populated, readable, and honest about being asserted."""

    def test_the_record_is_populated(self, capsys):
        payload = record(PLAIN, "--config", CONFIG, capsys=capsys)
        assert payload["controller"]["name"] == "ACME Finanz GmbH"
        assert payload["scope"]["provenance"]["from_mapping"] == 5

    def test_every_filled_value_is_labelled_as_an_assertion(self, capsys):
        payload = record(PLAIN, "--config", CONFIG, capsys=capsys)
        filled = [a for a in payload["activities"] if a["purpose"]["value"]]
        assert filled
        assert all(a["purpose"]["provenance"] == "mapping" for a in filled)

    def test_the_job_the_mapping_file_forgot_is_named_not_hidden(self, capsys):
        # The demo's whole argument: a register kept by hand drifts, and the
        # drift is invisible unless something says it out loud.
        payload = record(PLAIN, "--config", CONFIG, capsys=capsys)
        [missed] = [a for a in payload["activities"] if not a["purpose"]["value"]]
        assert missed["job"] == "acme.crm/consent-sync"
        assert any("1 of 6 activities has no purpose" in r for r in payload["reasons"])

    def test_and_the_record_still_does_not_claim_to_be_a_proof(self, capsys):
        assert record(PLAIN, "--config", CONFIG, capsys=capsys)["complete"] is False


class TestLineageThatDeclares:
    """Tier three: the same estate, materially richer, and it earns the mark."""

    def test_every_activity_stands_on_emitted_evidence(self, capsys):
        payload = record(DECLARED, "--config", CONFIG, capsys=capsys)
        assert payload["scope"]["provenance"] == {
            "evidenced": 6,
            "from_mapping": 0,
            "undeclared": 0,
        }
        assert payload["complete"] is True
        assert payload["reasons"] == []

    def test_the_facet_beats_the_mapping_file_that_is_still_present(self, capsys):
        # The same config is passed. A facet always wins, so nothing in the
        # record is asserted even though the file could have supplied it.
        payload = record(DECLARED, "--config", CONFIG, capsys=capsys)
        assert all(a["purpose"]["provenance"] == "facet" for a in payload["activities"])

    def test_the_difference_from_tier_two_is_material_not_cosmetic(self, capsys):
        asserted = record(PLAIN, "--config", CONFIG, capsys=capsys)
        evidenced = record(DECLARED, "--config", CONFIG, capsys=capsys)

        assert [a["job"] for a in asserted["activities"]] == [
            a["job"] for a in evidenced["activities"]
        ]
        assert asserted["scope"]["events"] == evidenced["scope"]["events"]
        # Same estate, same jobs, same events — and one is a proof.
        assert asserted["complete"] is not evidenced["complete"]

    def test_the_mark_reaches_the_terminal(self, capsys):
        assert main(["ropa", DECLARED, "--config", CONFIG]) == 0
        assert TOMBSTONE in capsys.readouterr().out


class TestTheWholeThingAsAnAuditorReceivesIt:
    def test_a_workbook_from_the_demo(self, tmp_path, capsys):
        out = tmp_path / "acme-ropa.xlsx"
        assert main(["ropa", DECLARED, "--config", CONFIG, "--out", str(out)]) == 0
        assert TOMBSTONE in capsys.readouterr().out

        book = load_workbook(out)
        sheet = book["Art. 30 record"]
        assert sheet["A2"].value == "ACME Finanz GmbH"
        assert "stands on emitted evidence" in sheet["A4"].value
        # Six activities under one header row.
        assert sheet.max_row == 12


class TestTheAssertionHistory:
    """`quality` against the same estate — and the half of it that is absence.

    The demo has two datasets with assertions and ten without, which is the
    proportion that makes the point: a report showing two green datasets out
    of twelve, without saying so, would be the invisible hole again.
    """

    def history(self, *argv, capsys) -> dict:
        assert main(["quality", *argv, "--format", "json"]) == 0
        return json.loads(capsys.readouterr().out)

    def test_it_finds_the_assertions_in_both_estates(self, capsys):
        # The two estates differ only in the Art. 30 facet, so the assertion
        # history must be identical across them.
        plain = self.history(PLAIN, "--config", CONFIG, capsys=capsys)
        declared = self.history(DECLARED, "--config", CONFIG, capsys=capsys)
        assert plain["datasets"] == declared["datasets"]
        assert plain["unchecked"] == declared["unchecked"]

    def test_two_datasets_are_checked_and_ten_are_not(self, capsys):
        payload = self.history(PLAIN, "--config", CONFIG, capsys=capsys)
        assert payload["scope"]["checked"] == 2
        assert payload["scope"]["unchecked"] == 10
        assert len(payload["unchecked"]) == 10

    def test_the_row_count_expectation_failed_twice_and_says_when(self, capsys):
        payload = self.history(PLAIN, "--config", CONFIG, capsys=capsys)
        [scored] = [d for d in payload["datasets"] if d["dataset"].endswith("transactions_scored")]
        [failing] = [e for e in scored["expectations"] if not e["holds"]]
        assert failing["assertion"] == "expect_table_row_count_to_be_between"
        assert failing["failures"] == 2
        assert failing["last_failure"].startswith("2026-06-10")

    def test_a_failing_expectation_does_not_withhold_the_mark(self, capsys):
        # It is withheld here, but for the unchecked datasets — not for the
        # failure. The mark says the history is complete, not that the data
        # is good.
        payload = self.history(PLAIN, "--config", CONFIG, capsys=capsys)
        assert payload["complete"] is False
        assert not any("failed" in r for r in payload["reasons"])
        assert any("carry no assertions" in r for r in payload["reasons"])

    def test_the_domain_filter_narrows_both_halves(self, capsys):
        payload = self.history(PLAIN, "--config", CONFIG, "--domain", "billing", capsys=capsys)
        assert [d["dataset"] for d in payload["datasets"]] == ["warehouse/billing_curated.invoices"]
        assert all("billing" in key for key in payload["unchecked"])

    def test_the_weekly_job_asserted_on_its_own_cadence(self, capsys):
        # Three weeks of Mondays, one of which failed before it could assert.
        payload = self.history(PLAIN, "--config", CONFIG, capsys=capsys)
        [invoices] = [d for d in payload["datasets"] if d["dataset"].endswith("invoices")]
        assert invoices["runs"] == 2


class TestTheProvenanceChain:
    """`provenance` against the same estate.

    The demo emits `sourceCodeLocation`, which dbt, Airflow and Spark all do,
    and nothing that reports a signature — because nothing standard does. So
    the chain names every commit and can still not claim authorisation, which
    is the honest state of the ecosystem rather than a gap in the demo.
    """

    def chain(self, dataset, *argv, capsys) -> dict:
        assert main(["provenance", PLAIN, "--dataset", dataset, *argv, "--format", "json"]) == 0
        return json.loads(capsys.readouterr().out)

    def test_it_walks_back_three_hops_to_the_source_tables(self, capsys):
        payload = self.chain("billing_curated.dunning_cases", capsys=capsys)
        assert [s["dataset"] for s in payload["steps"]] == [
            "warehouse/billing_curated.dunning_cases",
            "warehouse/billing_curated.invoices",
            "warehouse/billing_raw.orders",
            "warehouse/crm_curated.customers",
            "warehouse/crm_raw.accounts",
            "warehouse/crm_raw.contacts",
        ]

    def test_the_chain_crosses_a_domain_boundary(self, capsys):
        # billing's invoices read crm's customers. A provenance chain that
        # stopped at a domain edge would miss the interesting half.
        payload = self.chain("billing_curated.dunning_cases", capsys=capsys)
        jobs = {s["production"]["job"] for s in payload["steps"] if s["production"]}
        assert "acme.crm/customers-curated" in jobs
        assert "acme.billing/invoices-nightly" in jobs

    def test_every_produced_step_names_a_commit(self, capsys):
        payload = self.chain("billing_curated.dunning_cases", capsys=capsys)
        produced = [s for s in payload["steps"] if s["production"]]
        assert produced
        assert all(s["production"]["code"]["commit"] for s in produced)
        assert all(
            s["production"]["code"]["repository"].endswith("data-platform") for s in produced
        )

    def test_and_none_of_them_can_show_a_signature(self, capsys):
        payload = self.chain("billing_curated.dunning_cases", capsys=capsys)
        produced = [s for s in payload["steps"] if s["production"]]
        assert all(s["production"]["signed"] is None for s in produced)
        assert payload["complete"] is False
        assert any("unknown is not the same as unsigned" in r for r in payload["reasons"])

    def test_the_source_tables_are_ends_not_failures(self, capsys):
        payload = self.chain("billing_curated.dunning_cases", capsys=capsys)
        ends = [s for s in payload["steps"] if s["production"] is None]
        assert len(ends) == 3
        assert all("nothing in the window produced it" in s["ended"] for s in ends)
        assert payload["scope"]["ends_unproduced"] == 3

    def test_the_daily_job_wrote_the_dataset_many_times(self, capsys):
        # *The latest* must not read as *the only*.
        payload = self.chain("fraud_curated.transactions_scored", capsys=capsys)
        assert payload["steps"][0]["also_produced_by"] == 20

    def test_the_depth_limit_shows_up_as_its_own_kind_of_ending(self, capsys):
        payload = self.chain("billing_curated.dunning_cases", "--depth", "1", capsys=capsys)
        assert payload["scope"]["ends_at_depth"] >= 1
        assert any("depth limit of 1" in r for r in payload["reasons"])


ACTIVITIES = str(DEMO / "activities.yaml")


class TestTheDeployerViewOnTheDemo:
    """The transcript docs/evidence/deployer-view.md quotes, asserted.

    Same source, same config, one flag — and the view names the model version,
    the inputs, purpose and lawful basis, with the declared use case marked.
    """

    def deployer(self, source, *argv, capsys) -> dict:
        args = ["ropa", source, "--config", CONFIG, "--view", "deployer", *argv]
        assert main([*args, "--format", "json"]) == 0
        return json.loads(capsys.readouterr().out)

    def test_the_scoring_job_is_the_one_use_case(self, capsys):
        payload = self.deployer(DECLARED, capsys=capsys)
        assert [u["job"] for u in payload["use_cases"]] == ["acme.fraud/transactions-scored-daily"]
        assert payload["scope"]["activities"] == 6

    def test_it_names_both_model_versions_with_their_runs(self, capsys):
        [use_case] = self.deployer(DECLARED, capsys=capsys)["use_cases"]
        versions = {m["version"]: m["runs"] for m in use_case["model_versions"]}
        assert versions == {"2026-06-fraud-v3": 7, "2026-05-fraud-v2": 14}
        assert use_case["latest_run"]["model_version"] == "2026-06-fraud-v3"

    def test_it_names_the_inputs_the_latest_run_read(self, capsys):
        [use_case] = self.deployer(DECLARED, capsys=capsys)["use_cases"]
        assert use_case["latest_run"]["inputs"] == [
            "warehouse/fraud_raw.device_events",
            "warehouse/fraud_raw.transactions",
        ]

    def test_purpose_and_basis_are_evidenced_in_the_declared_lineage(self, capsys):
        payload = self.deployer(DECLARED, capsys=capsys)
        [use_case] = payload["use_cases"]
        assert use_case["purpose"] == {
            "value": "fraud-detection",
            "provenance": "facet",
            "unrecognised": False,
        }
        assert payload["complete"] is True

    def test_and_asserted_in_the_plain_lineage(self, capsys):
        payload = self.deployer(PLAIN, capsys=capsys)
        [use_case] = payload["use_cases"]
        assert use_case["purpose"]["provenance"] == "mapping"
        assert payload["complete"] is False

    def test_it_agrees_with_the_art_30_record_because_it_is_the_record(self, capsys):
        art30 = record(DECLARED, "--config", CONFIG, capsys=capsys)
        view = self.deployer(DECLARED, capsys=capsys)
        [activity] = [a for a in art30["activities"] if a["job"] == view["use_cases"][0]["job"]]
        assert activity["purpose"] == view["use_cases"][0]["purpose"]
        assert activity["legal_basis"] == view["use_cases"][0]["legal_basis"]

    def test_three_weeks_of_records_is_stated_as_a_span(self, capsys):
        [use_case] = self.deployer(DECLARED, capsys=capsys)["use_cases"]
        assert use_case["run_records"]["span_days"] == 20
        assert use_case["run_records"]["retention_policy"] is None

    def test_the_declared_use_case_is_marked_and_withholds_the_mark(self, capsys):
        payload = self.deployer(DECLARED, "--activities", ACTIVITIES, capsys=capsys)
        [entry] = payload["declared"]
        assert entry["name"] == "support-reply-drafts"
        assert entry["evidence"] == "declared"
        assert entry["legal_basis"]["provenance"] == "declared"
        assert payload["complete"] is False

    def test_the_model_version_tag_is_in_both_directories(self):
        # The two directories must still differ only in the processing facet.
        for directory in (DEMO / "lineage", DEMO / "lineage-declared"):
            body = (directory / "dbt.ndjson").read_text(encoding="utf-8")
            assert '"key": "model_version"' in body
