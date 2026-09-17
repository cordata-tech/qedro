"""The workbook.

The properties every format keeps — the scope statement, and provenance
surviving — are asserted for xlsx alongside the others in `test_render.py`.
What is here is what only a spreadsheet can get wrong: a verdict buried below
the table, a timestamp that silently lost its timezone, a dataset list that
stops without saying it stopped.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from qedro import TOMBSTONE, config, vocabulary
from qedro.events import parse_event
from qedro.ropa import build
from qedro.xlsx import AT, DATASETS_PER_CELL, HEADER_ROW, workbook

from .test_ropa import CONTROLLER, EVIDENCED, event

WORDS = vocabulary.load()


def book(events, cfg=CONTROLLER, **kw):
    record = build(events, config=config.parse(cfg), vocabulary=WORDS, **kw)
    return load_workbook(BytesIO(workbook(record)))


def cell(sheet, row, heading):
    return sheet.cell(row=row, column=AT[heading]).value


def _reading(count: int):
    """One evidenced job that reads *count* distinct datasets."""
    parsed = parse_event(
        {
            "eventType": "COMPLETE",
            "eventTime": "2026-03-01T10:00:00Z",
            "run": {"runId": "r1"},
            "job": {
                "namespace": "acme.fraud",
                "name": "scored",
                "facets": {"processing": dict(EVIDENCED)},
            },
            "inputs": [{"namespace": "wh", "name": f"raw.table_{n:03d}"} for n in range(count)],
        }
    )
    assert parsed is not None
    return parsed


@pytest.fixture
def evidenced():
    return book([event(facet=EVIDENCED)])


class TestTheVerdictIsWhereTheFileOpens:
    """A caveat below the table, or on a tab, is a caveat nobody reads.

    This is the tool's own thesis applied to its own output: a compliance
    artefact with holes is worse than none *because the holes are invisible*.
    """

    def test_it_sits_above_the_table_on_the_first_sheet(self):
        sheet = book([event()]).worksheets[0]
        assert sheet["A4"].value.startswith("This record does not claim to be a proof")
        assert sheet.cell(row=HEADER_ROW, column=1).value == "Activity"

    def test_and_says_so_when_the_record_does_stand_on_evidence(self, evidenced):
        sheet = evidenced.worksheets[0]
        assert "stands on emitted evidence" in sheet["A4"].value
        assert TOMBSTONE in sheet["A4"].value

    def test_the_count_of_reasons_is_not_a_plural_it_does_not_mean(self):
        # One reason, so `1 reason`. The same care as the CLI summary line.
        sheet = book([event(facet=EVIDENCED)], CONTROLLER + "domains: [fraud, marketing]\n")
        assert "1 reason on the Scope sheet" in sheet.worksheets[0]["A4"].value

    def test_the_reasons_themselves_are_on_the_scope_sheet(self):
        scope = book([event()])["Scope"]
        text = "\n".join(str(c.value) for row in scope.iter_rows() for c in row if c.value)
        assert "does not claim to be a proof" in text
        assert "no purpose or no legal basis" in text

    def test_a_missing_controller_is_stated_where_the_name_would_be(self):
        sheet = book([event(facet=EVIDENCED)], "domains: []\n").worksheets[0]
        assert sheet["A2"].value == "No controller is declared"


class TestTheTable:
    def test_two_sheets_named_for_what_they_hold(self, evidenced):
        assert evidenced.sheetnames == ["Art. 30 record", "Scope"]

    def test_one_row_per_activity_in_the_record_order(self):
        sheet = book([event(name="b"), event(name="a")]).worksheets[0]
        assert cell(sheet, HEADER_ROW + 1, "Activity") == "acme.fraud/a"
        assert cell(sheet, HEADER_ROW + 2, "Activity") == "acme.fraud/b"

    def test_the_header_is_frozen_and_filterable(self, evidenced):
        sheet = evidenced.worksheets[0]
        assert sheet.freeze_panes == f"A{HEADER_ROW + 1}"
        assert sheet.auto_filter.ref == f"A{HEADER_ROW}:L{HEADER_ROW + 1}"

    def test_an_evidenced_row_names_its_source(self, evidenced):
        sheet = evidenced.worksheets[0]
        assert cell(sheet, HEADER_ROW + 1, "Purpose") == "fraud-detection"
        assert cell(sheet, HEADER_ROW + 1, "Purpose source") == "emitted facet"

    def test_an_asserted_row_names_a_different_one(self):
        sheet = book(
            [event()], CONTROLLER + 'jobs:\n  "*": {purpose: p, legal_basis: consent}\n'
        ).worksheets[0]
        assert cell(sheet, HEADER_ROW + 1, "Purpose source") == "mapping file"
        assert cell(sheet, HEADER_ROW + 1, "Legal basis source") == "mapping file"

    def test_an_undeclared_row_says_nobody_said(self):
        sheet = book([event()]).worksheets[0]
        assert cell(sheet, HEADER_ROW + 1, "Purpose source") == "not declared"

    def test_the_notes_column_carries_why_a_row_is_not_evidenced(self):
        sheet = book([event()]).worksheets[0]
        assert "no purpose declared" in cell(sheet, HEADER_ROW + 1, "Notes")

    def test_a_record_with_no_activities_is_still_a_readable_file(self):
        sheet = book([]).worksheets[0]
        assert sheet["A1"].value == "Record of processing activities"
        assert sheet.cell(row=HEADER_ROW + 1, column=1).value is None


class TestTheThingsOnlyASpreadsheetGetsWrong:
    def test_a_timestamp_keeps_its_offset(self, evidenced):
        # Excel's date type has no timezone and openpyxl refuses an aware
        # datetime, so these are ISO strings. A real date cell would mean
        # dropping the offset from evidence about when something ran.
        seen = cell(evidenced.worksheets[0], HEADER_ROW + 1, "Last seen")
        assert isinstance(seen, str)
        assert seen == "2026-03-01T10:00:00+00:00"

    def test_a_truncated_dataset_list_says_that_it_was_truncated(self):
        # One job reading more datasets than fit in a cell. A list that simply
        # stops looks complete, which is the failure mode this whole tool is
        # about — so what was dropped is counted in the cell.
        sheet = book([_reading(DATASETS_PER_CELL + 5)]).worksheets[0]
        reads = cell(sheet, HEADER_ROW + 1, "Reads")
        assert len(reads.splitlines()) == DATASETS_PER_CELL + 1
        assert reads.splitlines()[-1] == "… and 5 more"

    def test_a_short_dataset_list_is_not_decorated(self, evidenced):
        reads = cell(evidenced.worksheets[0], HEADER_ROW + 1, "Reads")
        assert reads == "wh/raw.customers"

    def test_an_unrecognised_value_is_flagged_beside_the_value_not_inside_it(self):
        # The value column stays exactly what was declared, so a reader can
        # still sort and filter on it. The flag goes on the source cell.
        sheet = book([event(facet={"purpose": "p", "legal_basis": "vibes"})]).worksheets[0]
        assert cell(sheet, HEADER_ROW + 1, "Legal basis") == "vibes"
        assert "not in the vocabulary" in cell(sheet, HEADER_ROW + 1, "Legal basis source")

    def test_the_file_says_what_produced_it(self, evidenced):
        assert evidenced.properties.creator == "qedro"
        assert evidenced.properties.title == "Record of processing activities"

    def test_the_tint_marks_the_cell_the_problem_is_in(self):
        # Not evidenced tints the source; not in the vocabulary tints the
        # value. An emitted value can be sound evidence of a term nobody
        # agreed on, and those are different findings.
        sheet = book(
            [event(facet={"purpose": "p", "legal_basis": "vibes"})],
            CONTROLLER,
        ).worksheets[0]
        assert _tinted(sheet, HEADER_ROW + 1, "Legal basis")
        assert not _tinted(sheet, HEADER_ROW + 1, "Legal basis source")

        asserted = book(
            [event()], CONTROLLER + 'jobs:\n  "*": {purpose: p, legal_basis: consent}\n'
        ).worksheets[0]
        assert _tinted(asserted, HEADER_ROW + 1, "Purpose source")
        assert not _tinted(asserted, HEADER_ROW + 1, "Purpose")


def _tinted(sheet, row, heading) -> bool:
    fill = sheet.cell(row=row, column=AT[heading]).fill
    return bool(fill and fill.start_color and fill.start_color.rgb == "FFFDF0D5")


class TestTheScopeSheetCarriesEveryScopeLine:
    """The Art. 30 Scope sheet listed its own rows, and so never printed the
    declared count from #6. It now prints `scope.lines()` like every format."""

    def test_the_declared_count_reaches_the_workbook(self):
        from qedro import declared

        entries = declared.parse(
            {"activities": {"payroll": {"purpose": "payroll"}}}, origin="t", vocabulary=WORDS
        )
        record = build(
            [event(facet=EVIDENCED)],
            config=config.parse(CONTROLLER),
            vocabulary=WORDS,
            declared=entries,
        )
        sheet = load_workbook(BytesIO(workbook(record)))["Scope"]
        text = "\n".join(str(c.value) for row in sheet.iter_rows() for c in row if c.value)
        assert "1 activity declared with no lineage" in text
