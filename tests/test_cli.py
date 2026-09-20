"""The CLI, mostly to prove the mark rule holds end to end."""

import json

import pytest
from openpyxl import load_workbook

from qedro import TOMBSTONE, render
from qedro.__main__ import main

from .lineage_api import Backend, event

FIXTURES = "tests/fixtures/events"


def test_clean_read_earns_the_tombstone(capsys):
    assert main(["events", FIXTURES]) == 0
    out = capsys.readouterr().out
    assert "∎" in out
    assert "4 events" in out


def test_ascii_fallback_when_the_terminal_has_no_glyph(capsys):
    assert main(["events", FIXTURES, "--no-symbol"]) == 0
    out = capsys.readouterr().out
    assert "[complete]" in out
    assert "∎" not in out


def test_a_degraded_read_withholds_the_mark_and_says_why(tmp_path, capsys):
    good = json.dumps({"eventType": "COMPLETE", "run": {}, "job": {"namespace": "n", "name": "j"}})
    (tmp_path / "a.ndjson").write_text(f"{good}\n{{not json\n")
    assert main(["events", str(tmp_path)]) == 0
    cap = capsys.readouterr()
    assert "∎" not in cap.out
    assert "not usable OpenLineage events" in cap.err


def test_a_directory_that_yields_nothing_usable_fails(tmp_path, capsys):
    assert main(["events", str(tmp_path / "nope")]) == 1
    assert "does not exist" in capsys.readouterr().err


def test_no_command_prints_help_and_fails(capsys):
    assert main([]) == 1
    assert "usage:" in capsys.readouterr().err


def test_the_flag_works_on_either_side_of_the_subcommand(capsys):
    # `qedro --no-symbol events ./` and `qedro events ./ --no-symbol` are both
    # how people type it. argparse makes the second one hard, so it is tested.
    assert main(["--no-symbol", "events", FIXTURES]) == 0
    assert "[complete]" in capsys.readouterr().out

    assert main(["events", FIXTURES, "--no-symbol"]) == 0
    assert "[complete]" in capsys.readouterr().out


def test_the_env_var_suppresses_the_glyph_for_ci_logs(capsys, monkeypatch):
    monkeypatch.setenv("QEDRO_NO_SYMBOL", "1")
    assert main(["events", FIXTURES]) == 0
    out = capsys.readouterr().out
    assert "[complete]" in out and "∎" not in out


def test_the_summary_says_file_not_files_when_there_is_one(capsys):
    assert main(["events", FIXTURES]) == 0
    assert "1 file" in capsys.readouterr().out


def test_an_empty_window_does_not_earn_the_mark(capsys):
    # A clean read of nothing is not a proof of anything, and the difference
    # between "nothing happened" and "nothing was recorded" is exactly what a
    # reviewer needs flagged.
    assert main(["events", FIXTURES, "--since", "2030-01-01"]) == 0
    cap = capsys.readouterr()
    assert "0 events" in cap.out
    assert "∎" not in cap.out
    assert "nothing was recorded" in cap.err


def test_a_window_narrows_the_read(capsys):
    assert main(["events", FIXTURES, "--since", "2026-08-13", "--until", "2026-08-14"]) == 0
    assert "4 events" in capsys.readouterr().out


def test_an_unparseable_date_is_rejected_with_a_usable_message(capsys):
    with pytest.raises(SystemExit):
        main(["events", FIXTURES, "--since", "yesterday"])
    assert "ISO-8601" in capsys.readouterr().err


def test_an_address_that_is_not_a_directory_or_http_says_so(capsys):
    # `s3:/bucket does not exist` would send someone looking for a typo.
    assert main(["events", "s3://bucket/lineage"]) == 2
    assert "only a directory or an http:// address" in capsys.readouterr().err


class TestAgainstAnApi:
    """The acceptance test for source-neutrality: no AWS anywhere in this."""

    def test_reads_a_lineage_api_end_to_end(self, serve, capsys):
        base = serve(Backend(events=[event(f"j{i}") for i in range(3)]))
        assert main(["events", base]) == 0
        out = capsys.readouterr().out
        assert "3 events" in out and "1 page" in out
        assert "∎" in out

    def test_a_backend_that_is_down_fails_rather_than_reporting_nothing(self, serve, capsys):
        base = serve(Backend(status=503))
        assert main(["events", base]) == 1
        assert "HTTP 503" in capsys.readouterr().err


class TestRopa:
    """The projection, through the CLI people actually type."""

    def test_it_writes_a_record(self, tmp_path, capsys):
        events = tmp_path / "events"
        events.mkdir()
        (events / "e.json").write_text(
            json.dumps(
                {
                    "eventType": "COMPLETE",
                    "eventTime": "2026-03-01T10:00:00Z",
                    "run": {"runId": "r"},
                    "job": {
                        "namespace": "acme.fraud",
                        "name": "scored",
                        "facets": {
                            "processing": {
                                "purpose": "fraud-detection",
                                "legal_basis": "legitimate-interest",
                            }
                        },
                    },
                }
            ),
            encoding="utf-8",
        )
        config = tmp_path / "qedro.yaml"
        config.write_text("controller: ACME GmbH\n", encoding="utf-8")

        out = tmp_path / "ropa.md"
        code = main(
            [
                "ropa",
                str(events),
                "--config",
                str(config),
                "--format",
                "markdown",
                "--out",
                str(out),
            ]
        )
        assert code == 0
        assert "ACME GmbH" in out.read_text(encoding="utf-8")
        assert "wrote" in capsys.readouterr().out

    def test_the_mark_reaches_the_summary_line_when_earned(self, tmp_path, capsys):
        events, config = _estate(tmp_path, facet=True)
        out = tmp_path / "r.json"
        assert main(["ropa", str(events), "--config", str(config), "--out", str(out)]) == 0
        assert TOMBSTONE in capsys.readouterr().out

    def test_and_is_withheld_with_the_reason_on_stderr(self, tmp_path, capsys):
        events, config = _estate(tmp_path, facet=False)
        out = tmp_path / "r.json"
        assert main(["ropa", str(events), "--config", str(config), "--out", str(out)]) == 0
        captured = capsys.readouterr()
        assert TOMBSTONE not in captured.out
        # The file is still written. The run succeeded; it just does not claim
        # to be a proof.
        assert out.exists()
        assert "!" in captured.err

    def test_a_bad_config_is_a_sentence_not_a_traceback(self, tmp_path, capsys):
        events, _ = _estate(tmp_path, facet=True)
        bad = tmp_path / "bad.yaml"
        bad.write_text('jobs:\n  "*": {legalbasis: consent}\n', encoding="utf-8")
        assert main(["ropa", str(events), "--config", str(bad)]) == 2
        assert "unknown keys" in capsys.readouterr().err

    def test_a_named_config_that_does_not_exist_says_so(self, tmp_path, capsys):
        events, _ = _estate(tmp_path, facet=True)
        assert main(["ropa", str(events), "--config", str(tmp_path / "nope.yaml")]) == 2
        assert "no config file at" in capsys.readouterr().err

    def test_every_format_is_reachable(self, tmp_path, capsys):
        # Driven off the registry rather than a list written here, so a format
        # added later cannot be unreachable from the CLI and still pass.
        events, config = _estate(tmp_path, facet=True)
        for fmt in sorted(render.names()):
            argv = ["ropa", str(events), "--config", str(config), "--format", fmt]
            if render.is_binary(fmt):
                out = tmp_path / f"record.{fmt}"
                assert main([*argv, "--out", str(out)]) == 0
                assert out.stat().st_size > 0
            else:
                assert main(argv) == 0
                assert capsys.readouterr().out.strip()


def _estate(tmp_path, *, facet: bool):
    """A one-job directory and a config, with or without the emitted facet."""
    events = tmp_path / "events"
    events.mkdir(exist_ok=True)
    job = {"namespace": "acme.fraud", "name": "scored", "facets": {}}
    if facet:
        job["facets"]["processing"] = {
            "purpose": "fraud-detection",
            "legal_basis": "legitimate-interest",
        }
    (events / "e.json").write_text(
        json.dumps(
            {
                "eventType": "COMPLETE",
                "eventTime": "2026-03-01T10:00:00Z",
                "run": {"runId": "r"},
                "job": job,
            }
        ),
        encoding="utf-8",
    )
    config = tmp_path / "qedro.yaml"
    config.write_text("controller: ACME GmbH\n", encoding="utf-8")
    return events, config


class TestRopaAgainstAnApi:
    """The acceptance test for source-neutrality, on the projection this time.

    `qedro ropa` has to produce a record against a Marquez-compatible instance
    with **no AWS anywhere in the picture**. If this ever needs a LakeFormation
    client, a boto session or an AWS credential to pass, the tool has
    re-acquired the precondition that reading OpenLineage was supposed to
    remove.
    """

    def test_a_record_from_an_http_source(self, serve, tmp_path, capsys):
        emitted = event(job="scored")
        emitted["job"]["facets"] = {
            "processing": {
                "purpose": "fraud-detection",
                "legal_basis": "legitimate-interest",
            }
        }
        backend = Backend(events=[emitted])
        base = serve(backend)

        config = tmp_path / "qedro.yaml"
        config.write_text("controller: ACME GmbH\n", encoding="utf-8")

        assert main(["ropa", base, "--config", str(config), "--format", "json"]) == 0
        payload = json.loads(capsys.readouterr().out)

        assert payload["controller"]["name"] == "ACME GmbH"
        [activity] = payload["activities"]
        assert activity["purpose"]["provenance"] == "facet"
        assert payload["complete"] is True

        # Read-only by construction: the projection issued nothing but GETs.
        assert backend.methods == {"GET"}

    def test_the_scope_statement_names_the_instance(self, serve, tmp_path, capsys):
        # A record that does not say which instance it came from cannot be
        # reproduced by whoever receives it.
        base = serve(Backend(events=[event(job="scored")]))
        config = tmp_path / "qedro.yaml"
        config.write_text("controller: ACME GmbH\n", encoding="utf-8")

        assert main(["ropa", base, "--config", str(config), "--format", "json"]) == 0
        payload = json.loads(capsys.readouterr().out)
        assert payload["scope"]["source"] == base


class TestFormatFollowsTheFilename:
    def test_out_md_writes_markdown(self, tmp_path):
        events, config = _estate(tmp_path, facet=True)
        out = tmp_path / "ropa.md"
        assert main(["ropa", str(events), "--config", str(config), "--out", str(out)]) == 0
        assert out.read_text(encoding="utf-8").startswith("# Record of processing activities")

    def test_an_explicit_format_beats_the_filename(self, tmp_path):
        events, config = _estate(tmp_path, facet=True)
        out = tmp_path / "ropa.md"
        assert (
            main(
                [
                    "ropa",
                    str(events),
                    "--config",
                    str(config),
                    "--format",
                    "json",
                    "--out",
                    str(out),
                ]
            )
            == 0
        )
        json.loads(out.read_text(encoding="utf-8"))


class TestTheWorkbook:
    """xlsx is the one an auditor asks for, and the only output that is a file
    rather than something a terminal can show."""

    def test_out_xlsx_writes_a_workbook(self, tmp_path, capsys):
        events, config = _estate(tmp_path, facet=True)
        out = tmp_path / "ropa.xlsx"
        assert main(["ropa", str(events), "--config", str(config), "--out", str(out)]) == 0

        book = load_workbook(out)
        assert book.sheetnames == ["Art. 30 record", "Scope"]
        assert "wrote" in capsys.readouterr().out

    def test_asking_for_it_with_nowhere_to_put_it_is_a_sentence(self, tmp_path, capsys):
        events, config = _estate(tmp_path, facet=True)
        assert main(["ropa", str(events), "--config", str(config), "--format", "xlsx"]) == 2
        assert "add --out" in capsys.readouterr().err

    def test_that_failure_arrives_before_anything_is_read(self, capsys):
        # No source, no config, nothing that exists: the combination is
        # rejected on its own terms rather than after a long read.
        assert main(["ropa", "/nonexistent", "--format", "xlsx"]) == 2
        assert "add --out" in capsys.readouterr().err

    def test_the_mark_still_reaches_the_summary_line(self, tmp_path, capsys):
        events, config = _estate(tmp_path, facet=True)
        out = tmp_path / "ropa.xlsx"
        assert main(["ropa", str(events), "--config", str(config), "--out", str(out)]) == 0
        assert TOMBSTONE in capsys.readouterr().out


class TestQuality:
    """The assertion history, through the CLI people actually type."""

    def test_it_reports_what_was_checked_and_what_was_not(self, tmp_path, capsys):
        events, config = _asserting_estate(tmp_path)
        assert main(["quality", str(events), "--config", str(config), "--format", "json"]) == 0
        payload = json.loads(capsys.readouterr().out)

        assert [d["dataset"] for d in payload["datasets"]] == ["wh/scores"]
        assert payload["unchecked"] == ["wh/checked_nothing"]
        assert payload["complete"] is False

    def test_the_domain_filter_is_repeatable(self, tmp_path, capsys):
        events, config = _asserting_estate(tmp_path)
        assert (
            main(
                [
                    "quality",
                    str(events),
                    "--config",
                    str(config),
                    "--domain",
                    "fraud",
                    "--domain",
                    "billing",
                    "--format",
                    "json",
                ]
            )
            == 0
        )
        assert json.loads(capsys.readouterr().out)["scope"]["domains_seen"]

    def test_a_domain_nobody_emitted_for_is_an_empty_history_not_a_crash(self, tmp_path, capsys):
        events, config = _asserting_estate(tmp_path)
        argv = ["quality", str(events), "--config", str(config), "--domain", "nowhere"]
        assert main([*argv, "--format", "json"]) == 0
        assert json.loads(capsys.readouterr().out)["datasets"] == []

    def test_days_is_shorthand_for_since(self, tmp_path, capsys):
        # `--days 90` is how anybody asks for a quarter of history.
        events, config = _asserting_estate(tmp_path)
        assert main(["quality", str(events), "--config", str(config), "--days", "3650"]) == 0
        assert "assertion history" in capsys.readouterr().out

    def test_an_explicit_since_beats_days(self, tmp_path, capsys):
        # Two flags meaning the same thing must not silently disagree, and the
        # one the user typed is the one they meant.
        events, config = _asserting_estate(tmp_path)
        argv = ["quality", str(events), "--config", str(config), "--format", "json"]
        assert main([*argv, "--since", "2026-03-01", "--days", "1"]) == 0
        assert json.loads(capsys.readouterr().out)["scope"]["window"]["since"].startswith(
            "2026-03-01T00:00:00"
        )

    def test_every_format_is_reachable(self, tmp_path, capsys):
        events, config = _asserting_estate(tmp_path)
        for fmt in sorted(render.names()):
            argv = ["quality", str(events), "--config", str(config), "--format", fmt]
            if render.is_binary(fmt):
                out = tmp_path / f"history.{fmt}"
                assert main([*argv, "--out", str(out)]) == 0
                assert load_workbook(out).sheetnames == ["Assertions", "Not checked", "Scope"]
            else:
                assert main(argv) == 0
                assert capsys.readouterr().out.strip()

    def test_the_mark_is_earned_when_everything_in_view_was_checked(self, tmp_path, capsys):
        events, config = _asserting_estate(tmp_path, complete=True)
        assert main(["quality", str(events), "--config", str(config)]) == 0
        assert TOMBSTONE in capsys.readouterr().out


def _asserting_estate(tmp_path, *, complete: bool = False):
    """A job that asserts things, and (unless `complete`) one that does not."""
    events = tmp_path / "quality-events"
    events.mkdir(exist_ok=True)

    def event(job, reads, assertions=None):
        raw = {
            "eventType": "COMPLETE",
            "eventTime": "2026-03-01T10:00:00Z",
            "run": {"runId": job},
            "job": {"namespace": "acme.fraud", "name": job},
            "inputs": [{"namespace": "wh", "name": reads}],
        }
        if assertions:
            raw["inputs"][0]["inputFacets"] = {"dataQualityAssertions": {"assertions": assertions}}
        return raw

    rows = [
        event(
            "validate",
            "scores",
            [{"assertion": "expect_column_values_to_not_be_null", "column": "id", "success": True}],
        )
    ]
    if not complete:
        rows.append(event("load", "checked_nothing"))

    (events / "e.ndjson").write_text(
        "\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8"
    )
    config = tmp_path / "qedro.yaml"
    config.write_text("controller: ACME GmbH\n", encoding="utf-8")
    return events, config


class TestProvenance:
    """The chain, through the CLI people actually type."""

    def test_it_traces_a_dataset_by_bare_name(self, tmp_path, capsys):
        events, config = _chain_estate(tmp_path)
        argv = ["provenance", str(events), "--dataset", "final", "--config", str(config)]
        assert main([*argv, "--format", "json"]) == 0
        payload = json.loads(capsys.readouterr().out)

        assert payload["dataset"] == "wh/final"
        assert [s["dataset"] for s in payload["steps"]] == ["wh/final", "wh/mid", "wh/raw"]

    def test_a_name_that_matches_nothing_is_a_sentence(self, tmp_path, capsys):
        events, config = _chain_estate(tmp_path)
        argv = ["provenance", str(events), "--dataset", "nope", "--config", str(config)]
        assert main(argv) == 2
        assert "no dataset named" in capsys.readouterr().err

    def test_an_ambiguous_name_asks_rather_than_guesses(self, tmp_path, capsys):
        # Picking one would produce a chain for a dataset nobody asked about.
        events, config = _chain_estate(tmp_path, ambiguous=True)
        argv = ["provenance", str(events), "--dataset", "final", "--config", str(config)]
        assert main(argv) == 2
        err = capsys.readouterr().err
        assert "matches more than one dataset" in err
        assert "wh/final" in err and "lake/final" in err

    def test_the_depth_limit_is_reachable_and_reported(self, tmp_path, capsys):
        events, config = _chain_estate(tmp_path)
        argv = ["provenance", str(events), "--dataset", "final", "--config", str(config)]
        assert main([*argv, "--depth", "1", "--format", "json"]) == 0
        payload = json.loads(capsys.readouterr().out)
        assert payload["scope"]["depth_limit"] == 1
        assert payload["scope"]["ends_at_depth"] >= 1

    def test_every_format_is_reachable(self, tmp_path, capsys):
        events, config = _chain_estate(tmp_path)
        for fmt in sorted(render.names()):
            argv = [
                "provenance",
                str(events),
                "--dataset",
                "final",
                "--config",
                str(config),
                "--format",
                fmt,
            ]
            if render.is_binary(fmt):
                out = tmp_path / f"chain.{fmt}"
                assert main([*argv, "--out", str(out)]) == 0
                assert load_workbook(out).sheetnames == ["Chain", "Scope"]
            else:
                assert main(argv) == 0
                assert capsys.readouterr().out.strip()

    def test_the_mark_is_earned_when_the_whole_chain_is_signed(self, tmp_path, capsys):
        events, config = _chain_estate(tmp_path, signed=True, sources=False)
        argv = ["provenance", str(events), "--dataset", "final", "--config", str(config)]
        assert main(argv) == 0
        assert TOMBSTONE in capsys.readouterr().out


def _chain_estate(tmp_path, *, signed=None, ambiguous=False, sources=True):
    """final <- mid <- raw, with as much provenance evidence as asked for."""
    events = tmp_path / "chain-events"
    events.mkdir(exist_ok=True)

    def event(job, reads, writes, namespace="wh"):
        run_facets = {}
        if signed is not None:
            run_facets["cordata_provenance"] = {"descriptor_git_commit_signed": signed}
        return {
            "eventType": "COMPLETE",
            "eventTime": "2026-03-01T10:00:00Z",
            "run": {"runId": job, "facets": run_facets},
            "job": {
                "namespace": "acme.crm",
                "name": job,
                "facets": {
                    "sourceCodeLocation": {
                        "repoUrl": "https://github.com/acme/platform",
                        "version": f"commit-{job}",
                        "branch": "main",
                        "path": f"models/{job}.sql",
                    }
                },
            },
            "inputs": [{"namespace": "wh", "name": n} for n in reads],
            "outputs": [{"namespace": namespace, "name": n} for n in writes],
        }

    rows = [event("build-final", ["mid"], ["final"]), event("build-mid", ["raw"], ["mid"])]
    if sources:
        rows.append(event("ingest", [], ["raw"]))
    else:
        # No unproduced branch, so the only thing left to prove is the
        # signature — which is what the earned-mark test needs.
        rows = [event("build-final", [], ["final"])]
    if ambiguous:
        rows.append(event("other", [], ["final"], namespace="lake"))

    (events / "e.ndjson").write_text(
        "\n".join(json.dumps(r) for r in rows) + "\n", encoding="utf-8"
    )
    config = tmp_path / "qedro.yaml"
    config.write_text("controller: ACME GmbH\n", encoding="utf-8")
    return events, config


class TestTheDeployerView:
    """`qedro ropa --view deployer`, through the CLI people actually type."""

    def test_the_view_is_reachable_in_every_format(self, tmp_path, capsys):
        events, config = _estate(tmp_path, facet=True)
        for fmt in sorted(render.names()):
            argv = ["ropa", str(events), "--config", str(config), "--view", "deployer"]
            argv += ["--format", fmt]
            if render.is_binary(fmt):
                out = tmp_path / f"deployer.{fmt}"
                assert main([*argv, "--out", str(out)]) == 0
                assert load_workbook(out).sheetnames == ["AI use cases", "Scope"]
            else:
                assert main(argv) == 0
                assert capsys.readouterr().out.strip()

    def test_activities_are_read_in_the_art_30_view_too(self, tmp_path, capsys):
        # Refused there until cordata-tech/qedro#6 settled what they mean for
        # the record; now they are part of it.
        events, config = _estate(tmp_path, facet=True)
        doc = tmp_path / "activities.yaml"
        doc.write_text("activities:\n  payroll: {purpose: payroll, legal_basis: contract}\n")
        argv = ["ropa", str(events), "--config", str(config), "--activities", str(doc)]
        assert main([*argv, "--format", "json"]) == 0
        payload = json.loads(capsys.readouterr().out)
        assert [d["name"] for d in payload["declared"]] == ["payroll"]
        assert payload["complete"] is False

    def test_a_bad_activities_document_fails_before_the_events_are_read(self, tmp_path, capsys):
        bad = tmp_path / "activities.yaml"
        bad.write_text("activities:\n  x: {purpos: p}\n", encoding="utf-8")
        argv = ["ropa", str(tmp_path / "nowhere"), "--activities", str(bad)]
        assert main(argv) == 2
        assert "unknown keys purpos" in capsys.readouterr().err

    def test_a_bad_activities_document_is_a_sentence(self, tmp_path, capsys):
        events, config = _estate(tmp_path, facet=True)
        bad = tmp_path / "activities.yaml"
        bad.write_text("activities:\n  x: {purpose: p, legalbasis: contract}\n", encoding="utf-8")
        argv = ["ropa", str(events), "--config", str(config), "--view", "deployer"]
        assert main([*argv, "--activities", str(bad)]) == 2
        assert "unknown keys legalbasis" in capsys.readouterr().err


class TestTheDiffCommand:
    """`qedro diff` end to end — cordata-tech/qedro#14.

    The acceptance case in that issue is the demo estate: `demo/lineage` and
    `demo/lineage-declared` are identical apart from the `processing` facet, so
    comparing them is the provenance regression with nothing else moving.
    """

    @staticmethod
    def records(tmp_path, *, config="demo/qedro.yaml"):
        paths = []
        for name in ("lineage-declared", "lineage"):
            path = tmp_path / f"{name}.json"
            assert (
                main(
                    [
                        "ropa",
                        f"demo/{name}",
                        "--config",
                        config,
                        "--format",
                        "json",
                        "--out",
                        str(path),
                    ]
                )
                == 0
            )
            paths.append(str(path))
        return paths

    def test_it_reports_the_provenance_moving_from_facet_to_mapping(self, tmp_path, capsys):
        before, after = self.records(tmp_path)
        capsys.readouterr()

        assert main(["diff", before, after]) == 0
        out = capsys.readouterr().out
        assert "evidence lost: emitted facet → mapping file" in out
        # The value did not change, which is the whole point of the finding.
        assert "fraud-detection, unchanged" in out

    def test_it_reports_the_provenance_moving_from_mapping_to_facet(self, tmp_path, capsys):
        # The other direction, and the one #14's acceptance names: pipelines
        # that started declaring. A repair is reported as plainly as a loss, or
        # the command would only ever be read as bad news.
        declared, mapped = self.records(tmp_path)
        capsys.readouterr()

        assert main(["diff", mapped, declared]) == 0
        out = capsys.readouterr().out
        assert "now evidenced: mapping file → emitted facet" in out
        assert "0 of them a loss of evidence" in out

    def test_it_prints_no_mark_of_its_own(self, tmp_path, capsys):
        before, _ = self.records(tmp_path)
        capsys.readouterr()

        main(["diff", before, before])
        assert TOMBSTONE not in capsys.readouterr().out

    def test_it_reports_each_record_s_own_verdict(self, tmp_path, capsys):
        before, after = self.records(tmp_path)
        capsys.readouterr()

        main(["diff", before, after])
        out = capsys.readouterr().out
        assert "before        stands on its own evidence" in out
        assert "does not claim to be a proof" in out

    def test_cautions_reach_stderr_as_withheld_reasons_do(self, tmp_path, capsys):
        before, after = self.records(tmp_path)
        capsys.readouterr()

        main(["diff", before, after, "--out", str(tmp_path / "diff.md")])
        cap = capsys.readouterr()
        assert "wrote" in cap.out
        assert "12 findings, 12 a loss of evidence" in cap.out
        assert "different sources" in cap.err

    def test_the_format_is_inferred_from_the_filename(self, tmp_path):
        before, after = self.records(tmp_path)
        out = tmp_path / "diff.md"
        assert main(["diff", before, after, "--out", str(out)]) == 0
        assert out.read_text().startswith("# What changed between two records")

    def test_xlsx_without_out_fails_in_the_first_millisecond(self, tmp_path, capsys):
        before, after = self.records(tmp_path)
        capsys.readouterr()

        assert main(["diff", before, after, "--format", "xlsx"]) == 2
        assert "nowhere to put it" in capsys.readouterr().err

    def test_two_documents_of_different_shape_are_refused(self, tmp_path, capsys):
        before, _ = self.records(tmp_path)
        other = tmp_path / "quality.json"
        assert (
            main(
                [
                    "quality",
                    "demo/lineage",
                    "--config",
                    "demo/qedro.yaml",
                    "--format",
                    "json",
                    "--out",
                    str(other),
                ]
            )
            == 0
        )
        capsys.readouterr()

        assert main(["diff", before, str(other)]) == 2
        assert "different documents" in capsys.readouterr().err

    def test_a_file_that_is_not_a_record_is_a_sentence_not_a_traceback(self, tmp_path, capsys):
        path = tmp_path / "nope.json"
        path.write_text("{}")
        before, _ = self.records(tmp_path)
        capsys.readouterr()

        assert main(["diff", before, str(path)]) == 2
        assert "does not look like a record" in capsys.readouterr().err

    def test_comparing_a_record_with_itself_finds_nothing(self, tmp_path, capsys):
        before, _ = self.records(tmp_path)
        capsys.readouterr()

        assert main(["diff", before, before]) == 0
        out = capsys.readouterr().out
        assert "nothing in the Art. 30 content of these two records differs" in out
        assert "no findings" in out
